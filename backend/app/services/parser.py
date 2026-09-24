import io
import json
import logging
from typing import Dict, Any, Optional
from langchain_core.messages import HumanMessage
from app.agents.llm import llm_service, LLMResponseError

logger = logging.getLogger(__name__)

RESUME_PARSER_PROMPT = """你是一个专业的【简历信息抽取专家】。
请解析以下候选人简历文本，并提取关键结构化信息，严格返回纯 JSON 格式（无 markdown 包裹）：
```json
{{
  "name": "候选人姓名（简历中没有则填'候选人'）",
  "experience_years": 4,
  "skills": ["只列简历中真实出现的技能"],
  "projects": [
    {{
      "name": "项目名称（只能来自简历原文）",
      "role": "候选人在该项目中承担的角色",
      "tech_stack": ["该项目真实使用的技术"],
      "highlights": "候选人自己的可量化成果，不得编造"
    }}
  ],
  "education": "学历信息（只能来自简历原文）",
  "summary_profile": "基于简历内容的客观概括"
}}
```
硬性要求：示例仅用于说明字段结构，禁止沿用示例里的任何具体人名、公司、项目或技能；所有字段都必须来自简历原文，简历中没有的信息一律留空或写"未提供"，不要补全。
简历内容：
{resume_text}
"""

JD_PARSER_PROMPT = """你是一个专业的【岗位需求 (JD) 分析专家】。
请解析以下岗位描述 (Job Description) 文本，提取关键考核点与任职要求，严格返回纯 JSON 格式（无 markdown 包裹）：
```json
{{
  "title": "岗位名称（来自 JD 原文）",
  "level": "junior | senior | expert | director",
  "required_skills": ["JD 明确写出的硬性要求"],
  "preferred_skills": ["JD 写出的加分项，没有则为空数组"],
  "responsibilities": [
    "JD 原文中的职责，逐条提取"
  ],
  "interview_focus": [
    "由上述要求推导出的考察重点"
  ]
}}
```
硬性要求：示例仅用于说明字段结构，禁止沿用示例里的任何具体岗位、技能或职责；所有内容都必须来自 JD 原文，JD 没写的信息不要补全。
岗位描述内容：
{jd_text}
"""

class DocumentParserService:
    @staticmethod
    def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
        """Extract plain text from uploaded files (PDF, TXT, MD, DOCX)."""
        lower_name = filename.lower()
        
        # 1. PDF
        if lower_name.endswith(".pdf"):
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                pages_text = []
                for p in reader.pages:
                    txt = p.extract_text()
                    if txt:
                        pages_text.append(txt)
                return "\n".join(pages_text).strip()
            except Exception as e:
                logger.error(f"Failed to extract text from PDF: {e}")
                # Fallback to byte decode attempt
                return file_bytes.decode("utf-8", errors="ignore")

        # 2. TXT / Markdown
        if lower_name.endswith((".txt", ".md", ".json")):
            for enc in ("utf-8", "gbk", "gb2312", "latin1"):
                try:
                    return file_bytes.decode(enc)
                except UnicodeDecodeError:
                    continue
            return file_bytes.decode("utf-8", errors="ignore")

        # 3. DOCX（ZIP 容器，需用 python-docx 按段落提取，不能当纯文本解码）
        if lower_name.endswith(".docx"):
            try:
                import docx
                document = docx.Document(io.BytesIO(file_bytes))
                parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        row_text = " | ".join(
                            cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()
                        )
                        if row_text:
                            parts.append(row_text)
                return "\n".join(parts).strip()
            except Exception as e:
                logger.error(f"Failed to extract text from DOCX: {e}")
                return ""

        # 4. Others
        try:
            return file_bytes.decode("utf-8", errors="ignore")
        except Exception:
            return ""

    @staticmethod
    def _parse_json_response(content: str, what: str) -> Dict[str, Any]:
        """剥离 markdown 代码块后解析 JSON；不可解析时抛出可读的 LLMResponseError。"""
        text = (content or "").strip()
        if not text:
            raise LLMResponseError(f"模型没有返回{what}内容，请重试。")
        if "```json" in text:
            text = text.split("```json")[1].split("```")[0].strip()
        elif "```" in text:
            text = text.split("```")[1].split("```")[0].strip()
        try:
            parsed = json.loads(text)
        except json.JSONDecodeError as e:
            logger.error(f"{what} response is not valid JSON: {e}")
            raise LLMResponseError(f"模型返回的{what}结果不是合法 JSON，请重试。") from e
        if not isinstance(parsed, dict):
            raise LLMResponseError(f"模型返回的{what}结果格式不正确，请重试。")
        return parsed

    async def parse_resume(self, resume_text: str, llm_config: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """解析简历为结构化画像。模型不可用或返回异常时直接抛错，不返回示例数据。"""
        if not resume_text or not resume_text.strip():
            raise LLMResponseError("简历内容为空，请先粘贴简历文本或上传简历文件。")

        prompt = RESUME_PARSER_PROMPT.format(resume_text=resume_text)
        resp = await llm_service.invoke([HumanMessage(content=prompt)], llm_config=llm_config)
        return self._parse_json_response(resp.content, "简历画像")

    async def parse_jd(self, jd_text: str, llm_config: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        """解析 JD 为结构化岗位要求。模型不可用或返回异常时直接抛错，不返回示例数据。

        JD 为空是合法的"通用面试"场景：返回空结构，而不是编造一份岗位要求。
        """
        if not jd_text or not jd_text.strip():
            return {
                "title": "",
                "level": "",
                "required_skills": [],
                "preferred_skills": [],
                "responsibilities": [],
                "interview_focus": [],
            }

        prompt = JD_PARSER_PROMPT.format(jd_text=jd_text)
        resp = await llm_service.invoke([HumanMessage(content=prompt)], llm_config=llm_config)
        return self._parse_json_response(resp.content, "岗位要求")

parser_service = DocumentParserService()
